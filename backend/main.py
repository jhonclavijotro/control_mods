from fastapi import FastAPI, Depends, HTTPException, Query, UploadFile, File, Request, Response
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from sqlalchemy import or_
from pydantic import BaseModel, Field
from datetime import datetime, timedelta, timezone
import os
import shutil
import uuid
from typing import Optional, List

import hashlib
import secrets

try:
    from .database import engine, Base, get_db, SessionLocal, auto_migrate_db_schema
    from .models import Inverter, ModuleSlot, PowerModule, RepairLog, ReplacementLog, User
    from .solar_engine import calculate_module_metrics
except ImportError:
    import sys
    from pathlib import Path
    sys.path.append(str(Path(__file__).resolve().parent))
    from database import engine, Base, get_db, SessionLocal, auto_migrate_db_schema
    from models import Inverter, ModuleSlot, PowerModule, RepairLog, ReplacementLog, User
    from solar_engine import calculate_module_metrics

def get_utc_now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)

# Password Hashing and Session Verification Helpers
def hash_password(password: str) -> str:
    salt = secrets.token_hex(16)
    key = hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), salt.encode('utf-8'), 100000)
    return f"{salt}${key.hex()}"

def verify_password(password: str, password_hash: str) -> bool:
    try:
        salt, key_hex = password_hash.split("$")
        recalculated = hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), salt.encode('utf-8'), 100000)
        return secrets.compare_digest(recalculated.hex(), key_hex)
    except Exception:
        return False

# In-memory session token mapping: token -> user dict
SESSIONS = {}

# Create Database tables and auto-migrate schema
Base.metadata.create_all(bind=engine)
auto_migrate_db_schema()

# Configure Uploads Directory
DATA_DIR = os.path.dirname(os.path.abspath(engine.url.database)) if engine.url.database != ":memory:" else os.getcwd()
UPLOAD_DIR = os.path.join(DATA_DIR, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

# Read-Only Environment Configuration
READ_ONLY_MODE = os.getenv("READ_ONLY_MODE", "false").lower() in ["true", "1", "yes"]

app = FastAPI(
    title="Control de Módulos de Potencia - Granja Solar",
    description="API REST para seguimiento, mantenimiento y métricas de unidades de inversión y módulos de potencia.",
    version="1.2.0"
)

# Read-Only Middleware for Stakeholder Restricted Mode
@app.middleware("http")
async def read_only_middleware(request: Request, call_next):
    if READ_ONLY_MODE and request.method in ["POST", "PUT", "DELETE", "PATCH"]:
        if request.url.path in ["/api/auth/login", "/api/auth/logout", "/api/auth/register"]:
            return await call_next(request)
        return JSONResponse(
            status_code=403,
            content={"detail": "Modo de solo lectura activado (Stakeholders). Las operaciones de modificación están restringidas en este entorno."}
        )
    response = await call_next(request)
    return response

# Configure CORS safely
ALLOWED_ORIGINS = [o.strip() for o in os.getenv("ALLOWED_ORIGINS", "*").split(",") if o.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS if ALLOWED_ORIGINS != ["*"] else ["*"],
    allow_credentials=False if ALLOWED_ORIGINS == ["*"] else True,
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
    allow_headers=["*"],
)

app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")

ALLOWED_ATTACHMENT_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".docx", ".xlsx", ".txt"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB limit

@app.get("/api/config")
def get_app_config():
    """Returns runtime app configuration including read-only status."""
    return {
        "read_only_mode": READ_ONLY_MODE,
        "app_name": "Gestión de Módulos de Potencia - Granja Solar"
    }

@app.post("/api/upload")
async def upload_attachment(file: UploadFile = File(...)):
    """Upload an optional file attachment (document, photo, report) for maintenance logs."""
    if not file or not file.filename:
        raise HTTPException(status_code=400, detail="Ningún archivo seleccionado.")

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_ATTACHMENT_EXTENSIONS:
        allowed_str = ", ".join(sorted(ALLOWED_ATTACHMENT_EXTENSIONS))
        raise HTTPException(
            status_code=400, 
            detail=f"Tipo de archivo no permitido: '{ext}'. Formatos soportados: {allowed_str}"
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(status_code=413, detail="El archivo excede el tamaño máximo permitido (10 MB).")

    safe_basename = os.path.basename(file.filename)
    unique_name = f"{uuid.uuid4().hex}_{safe_basename}"
    dest_path = os.path.join(UPLOAD_DIR, unique_name)

    with open(dest_path, "wb") as buffer:
        buffer.write(content)

    web_url = f"/uploads/{unique_name}"
    return {
        "attachment_path": web_url,
        "attachment_name": safe_basename
    }

# Pydantic Request Models
class LoginRequest(BaseModel):
    username: str
    password: str

class RegisterStakeholderRequest(BaseModel):
    username: str
    password: str
    full_name: str

class CreateUserRequest(BaseModel):
    username: str
    password: str
    full_name: str
    role: str = "operator"

class StopRepairRequest(BaseModel):
    inverter_id: str
    slot_number: int
    stop_time: datetime
    reason: str
    diagnosis: Optional[str] = None
    attachment_path: Optional[str] = None
    attachment_name: Optional[str] = None

class RestartRepairRequest(BaseModel):
    repair_id: int
    restart_time: datetime
    diagnosis: Optional[str] = None
    attachment_path: Optional[str] = None
    attachment_name: Optional[str] = None

class RestartInverterRequest(BaseModel):
    inverter_id: str
    restart_time: datetime
    diagnosis: Optional[str] = None
    attachment_path: Optional[str] = None
    attachment_name: Optional[str] = None

class ModuleReplacementRequest(BaseModel):
    inverter_id: str
    slot_number: int
    new_serial: str
    reason: str
    timestamp: datetime = Field(default_factory=datetime.utcnow)
    performed_by: Optional[str] = "Técnico Solar"
    attachment_path: Optional[str] = None
    attachment_name: Optional[str] = None

class AddSpareModuleRequest(BaseModel):
    serial_number: str

REAL_AND_SYNTHETIC_SERIALS = {
    ("A1", 1): "30778700",
    ("A1", 2): "30778690",
    ("A1", 3): "30778681",
    ("A1", 4): "30778693",
    ("A1", 5): "30778687",
    ("A1", 6): "30778692",
    ("A2", 1): "30778672",
    ("A2", 2): "30778697",
    ("A2", 3): "30778661",
    ("A2", 4): "30778670",
    ("A2", 5): "30778673",
    ("A2", 6): "30773509",
    ("B1", 1): "30763827",
    ("B1", 2): "30773512",
    ("B1", 3): "30777967",
    ("B1", 4): "30777951",
    ("B1", 5): "30777962",
    ("B1", 6): "30778688",
    ("B2", 1): "30777965",
    ("B2", 2): "30778677",
    ("B2", 3): "30778676",
    ("B2", 4): "30778680",
    ("B2", 5): "30773498",
    ("B2", 6): "30778679",
    ("C1", 1): "30777977",
    ("C1", 2): "30778686",
    ("C1", 3): "30302105",
    ("C1", 4): "30778689",
    ("C1", 5): "SYNTH-C1-M5",
    ("C1", 6): "30778699",
    ("C2", 1): "30777955",
    ("C2", 2): "30778703",
    ("C2", 3): "30773510",
    ("C2", 4): "30777972",
    ("C2", 5): "30777976",
    ("C2", 6): "30778682",
    ("D1", 1): "30778678",
    ("D1", 2): "SYNTH-D1-M2",
    ("D1", 3): "SYNTH-D1-M3",
    ("D1", 4): "SYNTH-D1-M4",
    ("D1", 5): "30778684",
    ("D1", 6): "30778683",
    ("E1", 1): "30778685",
    ("E1", 2): "30778701",
    ("E1", 3): "30778691",
    ("E1", 4): "30777969",
}

def load_initial_serials():
    serials_map = dict(REAL_AND_SYNTHETIC_SERIALS)
    excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "BASE", "SERIALES.xlsx")
    if os.path.exists(excel_path):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(excel_path)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 3 and row[0] and row[1]:
                    inv = str(row[0]).strip().upper()
                    try:
                        slot = int(row[1])
                    except (ValueError, TypeError):
                        continue
                    raw_serial = row[2]
                    if raw_serial is not None and str(raw_serial).strip():
                        serials_map[(inv, slot)] = str(raw_serial).strip()
                    else:
                        serials_map[(inv, slot)] = f"SYNTH-{inv}-M{slot}"
        except Exception as e:
            print(f"Aviso al cargar Excel de seriales: {e}")
    return serials_map

# Seed initial farm data if empty
def seed_database_if_needed(db: Session):
    inverter_count = db.query(Inverter).count()
    if inverter_count > 0:
        return  # Database already initialized

    inverter_configs = [
        ("A1", "Unidad Inversora A1", 6),
        ("A2", "Unidad Inversora A2", 6),
        ("B1", "Unidad Inversora B1", 6),
        ("B2", "Unidad Inversora B2", 6),
        ("C1", "Unidad Inversora C1", 6),
        ("C2", "Unidad Inversora C2", 6),
        ("D1", "Unidad Inversora D1", 6),
        ("E1", "Unidad Inversora E1", 4),
    ]

    base_time = get_utc_now() - timedelta(days=30)
    serials_map = load_initial_serials()

    for inv_id, inv_name, max_mods in inverter_configs:
        inverter = Inverter(id=inv_id, name=inv_name, max_modules=max_mods)
        db.add(inverter)

        for slot_idx in range(1, max_mods + 1):
            serial = serials_map.get((inv_id, slot_idx), f"SYNTH-{inv_id}-M{slot_idx}")
            
            # Create Module Slot
            slot = ModuleSlot(
                inverter_id=inv_id,
                slot_number=slot_idx,
                current_serial=serial,
                installed_at=base_time
            )
            db.add(slot)

            # Create Power Module record
            pm = PowerModule(
                serial_number=serial,
                current_inverter_id=inv_id,
                current_slot_number=slot_idx,
                status="operating",
                registered_at=base_time,
                total_repairs=0
            )
            db.add(pm)

    # Add 6 Spare Modules in Inventory
    spare_serials = ["PM-SPARE-001", "PM-SPARE-002", "PM-SPARE-003", "PM-SPARE-004", "PM-SPARE-005", "PM-SPARE-006"]
    for sp_serial in spare_serials:
        sp_pm = PowerModule(
            serial_number=sp_serial,
            current_inverter_id=None,
            current_slot_number=None,
            status="spare",
            registered_at=base_time,
            total_repairs=0
        )
        db.add(sp_pm)

    db.commit()

def seed_default_users_if_needed(db: Session):
    admin_user = db.query(User).filter(User.username == "admin").first()
    if not admin_user:
        admin_user = User(
            username="admin",
            password_hash=hash_password("SolarisAdmin2026!"),
            full_name="Administrador del Sistema",
            role="admin",
            created_at=get_utc_now(),
            is_active=True
        )
        db.add(admin_user)

    stakeholder_user = db.query(User).filter(User.username == "stakeholder").first()
    if not stakeholder_user:
        stakeholder_user = User(
            username="stakeholder",
            password_hash=hash_password("SolarisViewer2026!"),
            full_name="Usuario Stakeholder",
            role="stakeholder",
            created_at=get_utc_now(),
            is_active=True
        )
        db.add(stakeholder_user)

    db.commit()

@app.on_event("startup")
def startup_event():
    db = SessionLocal()
    try:
        seed_default_users_if_needed(db)
        seed_database_if_needed(db)
    finally:
        db.close()

# API Endpoints

@app.get("/api/inverters")
def get_inverters(period: str = "month", db: Session = Depends(get_db)):
    """Return list of all 8 inverters with their slots and live status for a specified period filter."""
    inverters = db.query(Inverter).all()
    result = []
    
    now = get_utc_now()
    baseline_date = datetime(2026, 8, 13, 0, 0, 0)

    if period == "all_time":
        period_start = baseline_date
    elif period == "last_30":
        period_start = now - timedelta(days=30)
    elif period == "year":
        period_start = datetime(now.year, 1, 1, 0, 0, 0)
    else:
        # Default: current month accumulated
        period_start = datetime(now.year, now.month, 1, 0, 0, 0)

    for inv in inverters:
        slots_data = []
        for slot in inv.slots:
            pm = db.query(PowerModule).filter(PowerModule.serial_number == slot.current_serial).first()
            repair_logs = db.query(RepairLog).filter(RepairLog.serial_number == slot.current_serial).all()
            metrics = calculate_module_metrics(slot.installed_at, repair_logs, now, period_start=period_start)

            # Check if there is an active open repair
            active_repair = db.query(RepairLog).filter(
                RepairLog.serial_number == slot.current_serial,
                RepairLog.status == "open"
            ).first()

            effective_status = "in_repair" if active_repair else (pm.status if pm else "operating")

            slots_data.append({
                "slot_number": slot.slot_number,
                "current_serial": slot.current_serial,
                "installed_at": slot.installed_at.isoformat() if slot.installed_at else None,
                "status": effective_status,
                "metrics": metrics,
                "active_repair": {
                    "repair_id": active_repair.id,
                    "stop_time": active_repair.stop_time.isoformat(),
                    "reason": active_repair.reason,
                    "diagnosis": active_repair.diagnosis
                } if active_repair else None
            })

        # Sort slots by slot number
        slots_data.sort(key=lambda x: x["slot_number"])

        result.append({
            "id": inv.id,
            "name": inv.name,
            "max_modules": inv.max_modules,
            "slots": slots_data
        })

    return result

@app.get("/api/inverters/{inverter_id}")
def get_inverter_detail(inverter_id: str, db: Session = Depends(get_db)):
    """Return detailed information for a single inverter unit."""
    inv = db.query(Inverter).filter(Inverter.id == inverter_id.upper()).first()
    if not inv:
        raise HTTPException(status_code=404, detail="Unidad Inversora no encontrada")

    now = get_utc_now()
    slots_data = []
    for slot in inv.slots:
        pm = db.query(PowerModule).filter(PowerModule.serial_number == slot.current_serial).first()
        repair_logs = db.query(RepairLog).filter(RepairLog.serial_number == slot.current_serial).all()
        metrics = calculate_module_metrics(slot.installed_at, repair_logs, now)

        active_repair = db.query(RepairLog).filter(
            RepairLog.serial_number == slot.current_serial,
            RepairLog.status == "open"
        ).first()

        effective_status = "in_repair" if active_repair else (pm.status if pm else "operating")

        slots_data.append({
            "slot_number": slot.slot_number,
            "current_serial": slot.current_serial,
            "installed_at": slot.installed_at.isoformat() if slot.installed_at else None,
            "status": effective_status,
            "metrics": metrics,
            "active_repair": {
                "repair_id": active_repair.id,
                "stop_time": active_repair.stop_time.isoformat(),
                "reason": active_repair.reason,
                "diagnosis": active_repair.diagnosis
            } if active_repair else None
        })

    slots_data.sort(key=lambda x: x["slot_number"])

    return {
        "id": inv.id,
        "name": inv.name,
        "max_modules": inv.max_modules,
        "slots": slots_data
    }

@app.get("/api/dashboard")
def get_dashboard_summary(db: Session = Depends(get_db)):
    """Get global summary stats for the solar farm."""
    total_slots = db.query(ModuleSlot).count() or 46
    all_pms = db.query(PowerModule).filter(PowerModule.status != "spare", PowerModule.status != "retired").all()
    
    # Active open repairs override status
    open_repairs = db.query(RepairLog).filter(RepairLog.status == "open").all()
    open_repair_serials = set(r.serial_number for r in open_repairs)

    in_repair_count = len(open_repair_serials)
    operating_count = sum(1 for pm in all_pms if pm.serial_number not in open_repair_serials)
    spare_count = db.query(PowerModule).filter(PowerModule.status == "spare").count()

    # Calculate total solar operating hours farm-wide
    now = get_utc_now()
    slots = db.query(ModuleSlot).all()
    total_farm_hours = 0.0
    for slot in slots:
        repair_logs = db.query(RepairLog).filter(RepairLog.serial_number == slot.current_serial).all()
        metrics = calculate_module_metrics(slot.installed_at, repair_logs, now)
        total_farm_hours += metrics["net_operating_hours"]

    # Recent activity logs
    recent_repairs = db.query(RepairLog).order_by(RepairLog.stop_time.desc()).limit(5).all()
    recent_replacements = db.query(ReplacementLog).order_by(ReplacementLog.timestamp.desc()).limit(5).all()

    return {
        "read_only_mode": READ_ONLY_MODE,
        "total_active_slots": total_slots,
        "operating_count": operating_count,
        "in_repair_count": in_repair_count,
        "spare_count": spare_count,
        "total_farm_operating_hours": round(total_farm_hours, 1),
        "recent_repairs": [
            {
                "id": r.id,
                "serial": r.serial_number,
                "inverter_id": r.inverter_id,
                "slot_number": r.slot_number,
                "stop_time": r.stop_time.isoformat(),
                "restart_time": r.restart_time.isoformat() if r.restart_time else None,
                "reason": r.reason,
                "status": r.status
            } for r in recent_repairs
        ],
        "recent_replacements": [
            {
                "id": rep.id,
                "inverter_id": rep.inverter_id,
                "slot_number": rep.slot_number,
                "old_serial": rep.old_serial,
                "new_serial": rep.new_serial,
                "timestamp": rep.timestamp.isoformat(),
                "reason": rep.reason,
                "performed_by": rep.performed_by
            } for rep in recent_replacements
        ]
    }



@app.post("/api/repairs/stop")
def register_repair_stop(req: StopRepairRequest, db: Session = Depends(get_db)):
    """Register a stop event for a single power module slot OR for the entire inverter unit (slot_number=0)."""
    inv_id = req.inverter_id.upper()

    # If slot_number == 0, register a TOTAL INVERTER SHUTDOWN across all slots of that inverter
    if req.slot_number == 0:
        inv_slots = db.query(ModuleSlot).filter(ModuleSlot.inverter_id == inv_id).all()
        if not inv_slots:
            raise HTTPException(status_code=404, detail=f"Unidad Inversora {inv_id} no encontrada")

        stopped_count = 0
        for slot in inv_slots:
            if not slot.current_serial:
                continue

            existing_open = db.query(RepairLog).filter(
                RepairLog.serial_number == slot.current_serial,
                RepairLog.status == "open"
            ).first()

            if not existing_open:
                pm = db.query(PowerModule).filter(PowerModule.serial_number == slot.current_serial).first()
                if pm:
                    pm.status = "in_repair"
                    pm.total_repairs += 1

                repair = RepairLog(
                    serial_number=slot.current_serial,
                    inverter_id=inv_id,
                    slot_number=slot.slot_number,
                    stop_time=req.stop_time,
                    restart_time=None,
                    reason=f"[PARADA GENERAL DE INVERSOR {inv_id}] {req.reason}",
                    diagnosis=req.diagnosis or "Parada total de unidad inversora",
                    status="open",
                    attachment_path=req.attachment_path,
                    attachment_name=req.attachment_name
                )
                db.add(repair)
                stopped_count += 1

        db.commit()
        return {"message": f"Parada total de la Unidad Inversora {inv_id} registrada. {stopped_count} módulos detenidos."}

    # Individual slot stop
    slot = db.query(ModuleSlot).filter(
        ModuleSlot.inverter_id == inv_id,
        ModuleSlot.slot_number == req.slot_number
    ).first()

    if not slot or not slot.current_serial:
        raise HTTPException(status_code=404, detail="Slot o módulo no encontrado")

    pm = db.query(PowerModule).filter(PowerModule.serial_number == slot.current_serial).first()
    if pm:
        pm.status = "in_repair"
        pm.total_repairs += 1

    repair = RepairLog(
        serial_number=slot.current_serial,
        inverter_id=inv_id,
        slot_number=req.slot_number,
        stop_time=req.stop_time,
        restart_time=None,
        reason=req.reason,
        diagnosis=req.diagnosis,
        status="open",
        attachment_path=req.attachment_path,
        attachment_name=req.attachment_name
    )
    db.add(repair)
    db.commit()
    db.refresh(repair)

    return {"message": "Parada por reparación registrada exitosamente", "repair_id": repair.id}

@app.post("/api/repairs/restart-inverter")
def restart_inverter_all(req: RestartInverterRequest, db: Session = Depends(get_db)):
    """Restart all open repairs for a given inverter unit."""
    if not req.diagnosis or not req.diagnosis.strip():
        raise HTTPException(status_code=400, detail="Es obligatorio proporcionar un diagnóstico final o solución aplicada para reiniciar las paradas del inversor.")

    inv_id = req.inverter_id.upper()
    open_repairs = db.query(RepairLog).filter(
        RepairLog.inverter_id == inv_id,
        RepairLog.status == "open"
    ).all()

    if not open_repairs:
        raise HTTPException(status_code=404, detail=f"No hay paradas abiertas en la Unidad Inversora {inv_id}.")

    for r in open_repairs:
        if req.restart_time < r.stop_time:
            raise HTTPException(status_code=400, detail="La fecha de arranque no puede ser anterior a la fecha de parada")
        r.restart_time = req.restart_time
        r.status = "resolved"
        r.diagnosis = f"[ARRANQUE GENERAL] {req.diagnosis.strip()}"
        
        if req.attachment_path:
            r.attachment_path = req.attachment_path
            r.attachment_name = req.attachment_name

        pm = db.query(PowerModule).filter(PowerModule.serial_number == r.serial_number).first()
        if pm and pm.status == "in_repair":
            pm.status = "operating"

    db.commit()
    return {"message": f"Unidad Inversora {inv_id} restablecida exitosamente. Todos los módulos en marcha."}

@app.post("/api/repairs/restart")
def register_repair_restart(req: RestartRepairRequest, db: Session = Depends(get_db)):
    """Register restart of a module that was under repair."""
    if not req.diagnosis or not req.diagnosis.strip():
        raise HTTPException(status_code=400, detail="Es obligatorio proporcionar un diagnóstico final o solución aplicada para registrar el arranque.")

    repair = db.query(RepairLog).filter(RepairLog.id == req.repair_id).first()
    if not repair:
        raise HTTPException(status_code=404, detail="Registro de reparación no encontrado")

    if req.restart_time < repair.stop_time:
        raise HTTPException(status_code=400, detail="La fecha de arranque no puede ser anterior a la fecha de parada")

    repair.restart_time = req.restart_time
    repair.status = "resolved"
    repair.diagnosis = req.diagnosis.strip()

    if req.attachment_path:
        repair.attachment_path = req.attachment_path
        repair.attachment_name = req.attachment_name

    # Update module status back to operating if it is still installed
    pm = db.query(PowerModule).filter(PowerModule.serial_number == repair.serial_number).first()
    if pm and pm.status == "in_repair":
        pm.status = "operating"

    db.commit()
    return {"message": "Arranque de módulo registrado exitosamente"}

@app.post("/api/replacements")
def register_module_replacement(req: ModuleReplacementRequest, db: Session = Depends(get_db)):
    """Replace an existing module in an inverter slot with a new/spare module."""
    slot = db.query(ModuleSlot).filter(
        ModuleSlot.inverter_id == req.inverter_id.upper(),
        ModuleSlot.slot_number == req.slot_number
    ).first()

    if not slot:
        raise HTTPException(status_code=404, detail="Slot especificado no encontrado")

    old_serial = slot.current_serial

    # Verify new module exists or create if not in system
    new_pm = db.query(PowerModule).filter(PowerModule.serial_number == req.new_serial).first()
    if not new_pm:
        new_pm = PowerModule(
            serial_number=req.new_serial,
            current_inverter_id=req.inverter_id.upper(),
            current_slot_number=req.slot_number,
            status="operating",
            registered_at=req.timestamp,
            total_repairs=0
        )
        db.add(new_pm)
    else:
        new_pm.current_inverter_id = req.inverter_id.upper()
        new_pm.current_slot_number = req.slot_number
        new_pm.status = "operating"

    # Update old module status to 'retired' / 'in_repair'
    if old_serial:
        old_pm = db.query(PowerModule).filter(PowerModule.serial_number == old_serial).first()
        if old_pm:
            old_pm.current_inverter_id = None
            old_pm.current_slot_number = None
            old_pm.status = "retired"

    # Update slot current serial
    slot.current_serial = req.new_serial
    slot.installed_at = req.timestamp

    # Log replacement event
    rep_log = ReplacementLog(
        inverter_id=req.inverter_id.upper(),
        slot_number=req.slot_number,
        old_serial=old_serial or "NINGUNO",
        new_serial=req.new_serial,
        timestamp=req.timestamp,
        reason=req.reason,
        performed_by=req.performed_by or "Técnico Solar",
        attachment_path=req.attachment_path,
        attachment_name=req.attachment_name
    )
    db.add(rep_log)
    db.commit()

    return {"message": f"Módulo reemplazado en {req.inverter_id.upper()} slot {req.slot_number}. Nuevo serial: {req.new_serial}"}

@app.get("/api/spares")
def get_spares(db: Session = Depends(get_db)):
    """List available spare power modules."""
    spares = db.query(PowerModule).filter(PowerModule.status == "spare").all()
    return [{
        "serial_number": sp.serial_number,
        "registered_at": sp.registered_at.isoformat() if sp.registered_at else None,
        "status": sp.status,
        "total_repairs": sp.total_repairs
    } for sp in spares]

@app.post("/api/spares")
def add_spare(req: AddSpareModuleRequest, db: Session = Depends(get_db)):
    """Add a new spare power module to inventory."""
    existing = db.query(PowerModule).filter(PowerModule.serial_number == req.serial_number).first()
    if existing:
        raise HTTPException(status_code=400, detail="Un módulo con este número serial ya existe en el sistema.")

    sp = PowerModule(
        serial_number=req.serial_number,
        current_inverter_id=None,
        current_slot_number=None,
        status="spare",
        registered_at=get_utc_now(),
        total_repairs=0
    )
    db.add(sp)
    db.commit()
    return {"message": "Módulo de respaldo agregado exitosamente", "serial_number": req.serial_number}

@app.get("/api/logs")
def get_all_logs(
    inverter_id: Optional[str] = Query(None, description="Filtrar por ID de unidad inversora (ej: A1, B2)"),
    serial_number: Optional[str] = Query(None, description="Filtrar por número serial de módulo"),
    status: Optional[str] = Query(None, description="Filtrar reparaciones por estado: open, resolved, all"),
    search: Optional[str] = Query(None, description="Búsqueda por texto libre en motivos, diagnósticos y seriales"),
    db: Session = Depends(get_db)
):
    """Get audit logs for repairs and replacements with filtering support."""
    repair_query = db.query(RepairLog)
    replacement_query = db.query(ReplacementLog)

    if inverter_id and isinstance(inverter_id, str):
        inv_clean = inverter_id.strip().upper()
        if inv_clean:
            repair_query = repair_query.filter(RepairLog.inverter_id == inv_clean)
            replacement_query = replacement_query.filter(ReplacementLog.inverter_id == inv_clean)

    if serial_number and isinstance(serial_number, str):
        sn_clean = serial_number.strip()
        if sn_clean:
            pattern = f"%{sn_clean}%"
            repair_query = repair_query.filter(RepairLog.serial_number.ilike(pattern))
            replacement_query = replacement_query.filter(
                or_(
                    ReplacementLog.old_serial.ilike(pattern),
                    ReplacementLog.new_serial.ilike(pattern)
                )
            )

    if status and isinstance(status, str) and status.lower() != "all":
        st_clean = status.strip().lower()
        if st_clean:
            repair_query = repair_query.filter(RepairLog.status == st_clean)
            # Replacement logs do not have repair statuses (open/resolved)
            replacement_query = replacement_query.filter(ReplacementLog.id == -1)

    if search and isinstance(search, str):
        s_clean = search.strip()
        if s_clean:
            pattern = f"%{s_clean}%"
            repair_query = repair_query.filter(
                or_(
                    RepairLog.reason.ilike(pattern),
                    RepairLog.diagnosis.ilike(pattern),
                    RepairLog.serial_number.ilike(pattern),
                    RepairLog.inverter_id.ilike(pattern)
                )
            )
            replacement_query = replacement_query.filter(
                or_(
                    ReplacementLog.reason.ilike(pattern),
                    ReplacementLog.performed_by.ilike(pattern),
                    ReplacementLog.old_serial.ilike(pattern),
                    ReplacementLog.new_serial.ilike(pattern),
                    ReplacementLog.inverter_id.ilike(pattern)
                )
            )

    repairs = repair_query.order_by(RepairLog.stop_time.desc()).all()
    replacements = replacement_query.order_by(ReplacementLog.timestamp.desc()).all()

    return {
        "repairs": [
            {
                "id": r.id,
                "serial_number": r.serial_number,
                "inverter_id": r.inverter_id,
                "slot_number": r.slot_number,
                "stop_time": r.stop_time.isoformat(),
                "restart_time": r.restart_time.isoformat() if r.restart_time else None,
                "reason": r.reason,
                "diagnosis": r.diagnosis,
                "status": r.status,
                "attachment_path": r.attachment_path,
                "attachment_name": r.attachment_name
            } for r in repairs
        ],
        "replacements": [
            {
                "id": rep.id,
                "inverter_id": rep.inverter_id,
                "slot_number": rep.slot_number,
                "old_serial": rep.old_serial,
                "new_serial": rep.new_serial,
                "timestamp": rep.timestamp.isoformat(),
                "reason": rep.reason,
                "performed_by": rep.performed_by,
                "attachment_path": rep.attachment_path,
                "attachment_name": rep.attachment_name
            } for rep in replacements
        ]
    }

class EditSerialRequest(BaseModel):
    new_serial: str

class EditSlotInstallationDateRequest(BaseModel):
    installed_at: datetime

class EditRepairLogRequest(BaseModel):
    stop_time: datetime
    restart_time: Optional[datetime] = None
    reason: str
    diagnosis: Optional[str] = None
    attachment_path: Optional[str] = None
    attachment_name: Optional[str] = None

class EditReplacementLogRequest(BaseModel):
    timestamp: datetime
    reason: str
    performed_by: Optional[str] = "Técnico Solar"
    attachment_path: Optional[str] = None
    attachment_name: Optional[str] = None

@app.get("/api/modules")
def get_all_modules(db: Session = Depends(get_db)):
    """List all power modules in system with their status and metrics (batch queries optimized)."""
    modules = db.query(PowerModule).all()
    now = get_utc_now()

    # Pre-fetch slots into map: serial -> slot
    slots = db.query(ModuleSlot).all()
    slots_map = {s.current_serial: s for s in slots if s.current_serial}

    # Pre-fetch all repair logs grouped by serial_number
    all_repair_logs = db.query(RepairLog).all()
    repairs_map = {}
    active_repairs_map = {}
    for r in all_repair_logs:
        repairs_map.setdefault(r.serial_number, []).append(r)
        if r.status == "open":
            active_repairs_map[r.serial_number] = r

    result = []
    for pm in modules:
        slot = slots_map.get(pm.serial_number)
        repair_logs = repairs_map.get(pm.serial_number, [])
        installed_at = slot.installed_at if slot else pm.registered_at
        metrics = calculate_module_metrics(installed_at, repair_logs, now)

        active_repair = active_repairs_map.get(pm.serial_number)
        effective_status = "in_repair" if active_repair else pm.status

        result.append({
            "serial_number": pm.serial_number,
            "inverter_id": pm.current_inverter_id,
            "slot_number": pm.current_slot_number,
            "status": effective_status,
            "registered_at": pm.registered_at.isoformat() if pm.registered_at else None,
            "installed_at": installed_at.isoformat() if installed_at else None,
            "total_repairs": pm.total_repairs,
            "metrics": metrics
        })

    return result

@app.put("/api/modules/{old_serial}/edit-serial")
def edit_module_serial(old_serial: str, req: EditSerialRequest, db: Session = Depends(get_db)):
    """Edit the serial number of an existing power module across all system records."""
    new_serial = req.new_serial.strip()
    if not new_serial:
        raise HTTPException(status_code=400, detail="El nuevo número de serie no puede estar vacío.")

    if old_serial == new_serial:
        return {"message": "Sin cambios en el número de serie."}

    existing_new = db.query(PowerModule).filter(PowerModule.serial_number == new_serial).first()
    if existing_new:
        raise HTTPException(status_code=400, detail=f"Ya existe un módulo registrado con el número serial '{new_serial}'.")

    pm = db.query(PowerModule).filter(PowerModule.serial_number == old_serial).first()
    if not pm:
        raise HTTPException(status_code=404, detail=f"Módulo con serial '{old_serial}' no encontrado.")

    # Create new PowerModule with updated serial
    new_pm = PowerModule(
        serial_number=new_serial,
        current_inverter_id=pm.current_inverter_id,
        current_slot_number=pm.current_slot_number,
        status=pm.status,
        registered_at=pm.registered_at,
        total_repairs=pm.total_repairs
    )
    db.add(new_pm)

    # Update ModuleSlots
    slots = db.query(ModuleSlot).filter(ModuleSlot.current_serial == old_serial).all()
    for s in slots:
        s.current_serial = new_serial

    # Update RepairLogs
    repairs = db.query(RepairLog).filter(RepairLog.serial_number == old_serial).all()
    for r in repairs:
        r.serial_number = new_serial

    # Update ReplacementLogs
    reps_old = db.query(ReplacementLog).filter(ReplacementLog.old_serial == old_serial).all()
    for ro in reps_old:
        ro.old_serial = new_serial

    reps_new = db.query(ReplacementLog).filter(ReplacementLog.new_serial == old_serial).all()
    for rn in reps_new:
        rn.new_serial = new_serial

    db.delete(pm)
    db.commit()

    return {"message": f"Serial actualizado exitosamente de '{old_serial}' a '{new_serial}'."}

@app.delete("/api/modules/{serial_number}")
def delete_module(serial_number: str, db: Session = Depends(get_db)):
    """Delete a power module from inventory."""
    pm = db.query(PowerModule).filter(PowerModule.serial_number == serial_number).first()
    if not pm:
        raise HTTPException(status_code=404, detail=f"Módulo con serial '{serial_number}' no encontrado.")

    # Unassign from any slot if installed
    slots = db.query(ModuleSlot).filter(ModuleSlot.current_serial == serial_number).all()
    for s in slots:
        s.current_serial = None

    db.delete(pm)
    db.commit()
    return {"message": f"Módulo '{serial_number}' eliminado exitosamente del inventario."}

@app.put("/api/slots/{inverter_id}/{slot_number}/installed-at")
def edit_slot_installation_date(
    inverter_id: str,
    slot_number: int,
    req: EditSlotInstallationDateRequest,
    db: Session = Depends(get_db)
):
    """Edit the installation date and time (installed_at) for a specific module slot."""
    slot = db.query(ModuleSlot).filter(
        ModuleSlot.inverter_id == inverter_id.upper(),
        ModuleSlot.slot_number == slot_number
    ).first()

    if not slot:
        raise HTTPException(status_code=404, detail=f"Slot {slot_number} de la unidad {inverter_id.upper()} no encontrado.")

    slot.installed_at = req.installed_at
    
    if slot.current_serial:
        pm = db.query(PowerModule).filter(PowerModule.serial_number == slot.current_serial).first()
        if pm:
            pm.registered_at = req.installed_at

    db.commit()
    return {
        "message": f"Fecha y hora de instalación actualizada exitosamente para {inverter_id.upper()} slot {slot_number}.",
        "installed_at": slot.installed_at.isoformat()
    }

@app.put("/api/repairs/{repair_id}")
def edit_repair_log(repair_id: int, req: EditRepairLogRequest, db: Session = Depends(get_db)):
    """Edit details of an existing repair log event for human error correction."""
    repair = db.query(RepairLog).filter(RepairLog.id == repair_id).first()
    if not repair:
        raise HTTPException(status_code=404, detail="Registro de reparación no encontrado.")

    if req.restart_time and req.restart_time < req.stop_time:
        raise HTTPException(status_code=400, detail="La fecha de arranque no puede ser anterior a la fecha de parada.")

    final_diag = req.diagnosis.strip() if req.diagnosis else (repair.diagnosis or "")
    if req.restart_time and not final_diag.strip():
        raise HTTPException(status_code=400, detail="No se puede registrar la fecha de arranque o resolver la parada sin proporcionar un diagnóstico final o solución aplicada.")

    repair.stop_time = req.stop_time
    repair.restart_time = req.restart_time
    repair.reason = req.reason
    if req.diagnosis is not None:
        repair.diagnosis = req.diagnosis.strip()

    if req.attachment_path:
        repair.attachment_path = req.attachment_path
        repair.attachment_name = req.attachment_name

    if req.restart_time:
        repair.status = "resolved"
        pm = db.query(PowerModule).filter(PowerModule.serial_number == repair.serial_number).first()
        if pm and pm.status == "in_repair":
            pm.status = "operating"
    else:
        repair.status = "open"
        pm = db.query(PowerModule).filter(PowerModule.serial_number == repair.serial_number).first()
        if pm and pm.status == "operating":
            pm.status = "in_repair"

    db.commit()
    return {"message": "Registro de reparación actualizado exitosamente."}

@app.put("/api/replacements/{replacement_id}")
def edit_replacement_log(replacement_id: int, req: EditReplacementLogRequest, db: Session = Depends(get_db)):
    """Edit details of an existing replacement log event for human error correction."""
    rep = db.query(ReplacementLog).filter(ReplacementLog.id == replacement_id).first()
    if not rep:
        raise HTTPException(status_code=404, detail="Registro de reemplazo no encontrado.")

    rep.timestamp = req.timestamp
    rep.reason = req.reason
    if req.performed_by:
        rep.performed_by = req.performed_by

    if req.attachment_path:
        rep.attachment_path = req.attachment_path
        rep.attachment_name = req.attachment_name

    db.commit()
    return {"message": "Registro de reemplazo actualizado exitosamente."}

@app.post("/api/seed/clean")
def clean_database_for_production(db: Session = Depends(get_db)):
    """Clear all repairs, replacements, and reset all operating hours to 0 while keeping current serial numbers intact."""
    if READ_ONLY_MODE:
        raise HTTPException(status_code=403, detail="Operación restringida en Modo Solo Lectura.")

    db.query(RepairLog).delete()
    db.query(ReplacementLog).delete()

    now = get_utc_now()

    # Reset slots installed_at to now so operating hours start at 0.0 hrs
    slots = db.query(ModuleSlot).all()
    for s in slots:
        s.installed_at = now

    # Reset modules status & repair count without altering serial_number
    modules = db.query(PowerModule).all()
    for pm in modules:
        pm.total_repairs = 0
        if pm.status != "spare" and pm.status != "retired":
            pm.status = "operating"
        pm.registered_at = now

    db.commit()
    return {"message": "Horas de operación restablecidas a CERO y reparaciones eliminadas. Los números seriales se han conservado."}

@app.post("/api/seed/reset")
def reset_database(db: Session = Depends(get_db)):
    """Reset and re-seed database with default solar farm configuration."""
    if READ_ONLY_MODE:
        raise HTTPException(status_code=403, detail="Operación restringida en Modo Solo Lectura.")

    db.query(RepairLog).delete()
    db.query(ReplacementLog).delete()
    db.query(ModuleSlot).delete()
    db.query(PowerModule).delete()
    db.query(Inverter).delete()
    db.commit()

    seed_default_users_if_needed(db)
    seed_database_if_needed(db)
    return {"message": "Base de datos reiniciada con datos por defecto de la granja solar."}

@app.get("/api/export/excel")
def export_modules_excel(db: Session = Depends(get_db)):
    """Export complete power modules list (installed & stock) as Excel-compatible CSV file."""
    import csv
    import io

    now = get_utc_now()
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')

    writer.writerow([
        "Serial",
        "Ubicación / Estado",
        "Día de Instalación",
        "Día de Retiro",
        "Tiempo Total de Operación (hrs)",
        "Porcentaje de Inoperancia (%)",
        "Número de Revisiones"
    ])

    modules = db.query(PowerModule).all()
    for pm in modules:
        slot = db.query(ModuleSlot).filter(ModuleSlot.current_serial == pm.serial_number).first()
        repair_logs = db.query(RepairLog).filter(RepairLog.serial_number == pm.serial_number).all()

        if slot:
            inv_id = slot.inverter_id
            location = f"Inversor {inv_id} - Slot {slot.slot_number}"
            install_date = slot.installed_at.strftime("%d/%m/%Y %I:%M %p") if slot.installed_at else "13/08/2026 07:00 AM"
            retired_date = "Presente"
            metrics = calculate_module_metrics(slot.installed_at, repair_logs, now)
            op_hours = f"{metrics['net_operating_hours']:.1f}"
            inop_pct = f"{100.0 - metrics['uptime_percent']:.1f}%"
            revisions = len(repair_logs)
        else:
            location = "Almacén Central (Stock)" if pm.status == "spare" else "Retirado"
            install_date = "N/A (Stock)" if pm.status == "spare" else (pm.registered_at.strftime("%d/%m/%Y") if pm.registered_at else "N/A")
            retired_date = "Presente" if pm.status == "spare" else "Retirado"
            op_hours = "0.0"
            inop_pct = "0.0%"
            revisions = pm.total_repairs or len(repair_logs)

        writer.writerow([
            pm.serial_number,
            location,
            install_date,
            retired_date,
            op_hours,
            inop_pct,
            revisions
        ])

    csv_text = output.getvalue()
    return Response(
        content=csv_text.encode("utf-8-sig"),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="Reporte_Modulos_Solaris_Control.csv"'}
    )

# Authentication and User Management Endpoints

@app.post("/api/auth/login")
def login(req: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == req.username.strip().lower()).first()
    if not user or not user.is_active or not verify_password(req.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Usuario o contraseña incorrectos.")

    # Restrict login on public Ngrok / Stakeholder instance (READ_ONLY_MODE=true) to stakeholder role only
    if READ_ONLY_MODE and user.role != "stakeholder":
        raise HTTPException(
            status_code=403, 
            detail="Acceso restringido: En el portal público de Stakeholders únicamente se permite el ingreso a usuarios con rol Stakeholder."
        )

    token = secrets.token_hex(32)
    user_info = {
        "id": user.id,
        "username": user.username,
        "full_name": user.full_name,
        "role": user.role
    }
    SESSIONS[token] = user_info

    return {
        "message": "Inicio de sesión exitoso",
        "token": token,
        "user": user_info
    }

@app.post("/api/auth/register")
def register_stakeholder(req: RegisterStakeholderRequest, db: Session = Depends(get_db)):
    username_clean = req.username.strip().lower()
    if not username_clean or len(username_clean) < 3:
        raise HTTPException(status_code=400, detail="El nombre de usuario debe tener al menos 3 caracteres.")

    if len(req.password.strip()) < 4:
        raise HTTPException(status_code=400, detail="La contraseña debe tener al menos 4 caracteres.")

    existing = db.query(User).filter(User.username == username_clean).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"El nombre de usuario '{username_clean}' ya se encuentra registrado.")

    new_user = User(
        username=username_clean,
        password_hash=hash_password(req.password.strip()),
        full_name=req.full_name.strip() or username_clean,
        role="stakeholder",
        created_at=get_utc_now(),
        is_active=True
    )
    db.add(new_user)
    db.commit()

    # Auto login upon registration
    token = secrets.token_hex(32)
    user_info = {
        "id": new_user.id,
        "username": new_user.username,
        "full_name": new_user.full_name,
        "role": new_user.role
    }
    SESSIONS[token] = user_info

    return {
        "message": "Registro de Stakeholder exitoso",
        "token": token,
        "user": user_info
    }

@app.post("/api/auth/logout")
def logout(request: Request):
    auth_header = request.headers.get("Authorization") or ""
    token = auth_header.replace("Bearer ", "").strip()
    if token in SESSIONS:
        del SESSIONS[token]
    return {"message": "Sesión cerrada exitosamente."}

@app.get("/api/auth/me")
def get_current_user_profile(request: Request):
    auth_header = request.headers.get("Authorization") or ""
    token = auth_header.replace("Bearer ", "").strip()
    if not token or token not in SESSIONS:
        raise HTTPException(status_code=401, detail="Sesión no válida o expirada.")
    return SESSIONS[token]

@app.get("/api/users")
def get_users(request: Request, db: Session = Depends(get_db)):
    auth_header = request.headers.get("Authorization") or ""
    token = auth_header.replace("Bearer ", "").strip()
    user_info = SESSIONS.get(token)
    if not user_info or user_info.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Acceso permitido únicamente para administradores.")

    users = db.query(User).all()
    return [{
        "id": u.id,
        "username": u.username,
        "full_name": u.full_name,
        "role": u.role,
        "is_active": u.is_active,
        "created_at": u.created_at.isoformat() if u.created_at else None
    } for u in users]

@app.post("/api/users")
def create_user(req: CreateUserRequest, request: Request, db: Session = Depends(get_db)):
    auth_header = request.headers.get("Authorization") or ""
    token = auth_header.replace("Bearer ", "").strip()
    user_info = SESSIONS.get(token)
    if not user_info or user_info.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Acceso permitido únicamente para administradores.")

    username_clean = req.username.strip().lower()
    if not username_clean or len(username_clean) < 3:
        raise HTTPException(status_code=400, detail="El nombre de usuario debe tener al menos 3 caracteres.")

    existing = db.query(User).filter(User.username == username_clean).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"El nombre de usuario '{username_clean}' ya está registrado.")

    new_user = User(
        username=username_clean,
        password_hash=hash_password(req.password.strip()),
        full_name=req.full_name.strip(),
        role=req.role.lower() if req.role in ["admin", "operator", "stakeholder"] else "operator",
        created_at=get_utc_now(),
        is_active=True
    )
    db.add(new_user)
    db.commit()
    return {"message": f"Usuario '{username_clean}' creado exitosamente."}

@app.delete("/api/users/{user_id}")
def delete_user(user_id: int, request: Request, db: Session = Depends(get_db)):
    auth_header = request.headers.get("Authorization") or ""
    token = auth_header.replace("Bearer ", "").strip()
    user_info = SESSIONS.get(token)
    if not user_info or user_info.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Acceso permitido únicamente para administradores.")

    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Usuario no encontrado.")

    if user.username == "admin" or user.id == user_info.get("id"):
        raise HTTPException(status_code=400, detail="No es posible eliminar la cuenta del administrador actual.")

    db.delete(user)
    db.commit()
    return {"message": f"Usuario '{user.username}' eliminado exitosamente."}

@app.get("/api/export/xlsx")
def export_modules_xlsx(db: Session = Depends(get_db)):
    """Export complete solar farm data (modules, repairs, replacements) as a multi-sheet Excel .xlsx file."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import io

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    border_side = Side(style="thin", color="CBD5E1")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Sheet 1: Módulos de Potencia
    ws1 = wb.create_sheet(title="Módulos de Potencia")
    ws1.append(["SOLARIS CONTROL - REPORTES DE MÓDULOS DE POTENCIA"])
    ws1.cell(row=1, column=1).font = title_font
    ws1.append([])

    headers_ws1 = [
        "Serial Number", "Ubicación / Estado", "Fecha Instalación", 
        "Horas Potenciales", "Horas Inoperancia", "Horas Operación Neta", 
        "Disponibilidad (%)", "MTBF (hrs)", "Total Revisiones"
    ]
    ws1.append(headers_ws1)

    now = get_utc_now()
    modules = db.query(PowerModule).all()
    for pm in modules:
        slot = db.query(ModuleSlot).filter(ModuleSlot.current_serial == pm.serial_number).first()
        repair_logs = db.query(RepairLog).filter(RepairLog.serial_number == pm.serial_number).all()

        if slot:
            location = f"Inversor {slot.inverter_id} - Slot {slot.slot_number}"
            install_date = slot.installed_at.strftime("%d/%m/%Y %H:%M") if slot.installed_at else "N/A"
            metrics = calculate_module_metrics(slot.installed_at, repair_logs, now)
        else:
            location = "Almacén Central (Stock)" if pm.status == "spare" else "Retirado"
            install_date = "N/A (Stock)" if pm.status == "spare" else (pm.registered_at.strftime("%d/%m/%Y") if pm.registered_at else "N/A")
            metrics = {
                "total_potential_solar_hours": 0.0,
                "solar_downtime_hours": 0.0,
                "net_operating_hours": 0.0,
                "uptime_percent": 100.0,
                "mtbf_hours": 0.0
            }

        ws1.append([
            pm.serial_number,
            location,
            install_date,
            metrics["total_potential_solar_hours"],
            metrics["solar_downtime_hours"],
            metrics["net_operating_hours"],
            f"{metrics['uptime_percent']:.1f}%",
            metrics["mtbf_hours"],
            len(repair_logs)
        ])

    # Sheet 2: Historial de Reparaciones
    ws2 = wb.create_sheet(title="Historial Reparaciones")
    ws2.append(["HISTORIAL DE PARADAS Y REPARACIONES DE MÓDULOS"])
    ws2.cell(row=1, column=1).font = title_font
    ws2.append([])

    headers_ws2 = ["ID", "Serial Módulo", "Unidad Inversora", "Slot", "Fecha Parada", "Fecha Arranque", "Motivo", "Diagnóstico", "Estado"]
    ws2.append(headers_ws2)

    repairs = db.query(RepairLog).order_by(RepairLog.stop_time.desc()).all()
    for r in repairs:
        ws2.append([
            r.id,
            r.serial_number,
            r.inverter_id,
            r.slot_number,
            r.stop_time.strftime("%d/%m/%Y %H:%M") if r.stop_time else "",
            r.restart_time.strftime("%d/%m/%Y %H:%M") if r.restart_time else "EN REPARACIÓN",
            r.reason,
            r.diagnosis or "Pendiente",
            "ABIERTA" if r.status == "open" else "RESUELTA"
        ])

    # Sheet 3: Historial de Reemplazos
    ws3 = wb.create_sheet(title="Historial Reemplazos")
    ws3.append(["REGISTRO DE REEMPLAZOS DE MÓDULOS EN SLOTS"])
    ws3.cell(row=1, column=1).font = title_font
    ws3.append([])

    headers_ws3 = ["ID", "Unidad Inversora", "Slot", "Serial Antiguo", "Serial Nuevo", "Fecha Reemplazo", "Motivo", "Realizado Por"]
    ws3.append(headers_ws3)

    replacements = db.query(ReplacementLog).order_by(ReplacementLog.timestamp.desc()).all()
    for rep in replacements:
        ws3.append([
            rep.id,
            rep.inverter_id,
            rep.slot_number,
            rep.old_serial,
            rep.new_serial,
            rep.timestamp.strftime("%d/%m/%Y %H:%M") if rep.timestamp else "",
            rep.reason,
            rep.performed_by or "Técnico Solar"
        ])

    # Styling and column width calculation
    for ws in [ws1, ws2, ws3]:
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = thin_border
                if cell.row > 3 and isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Solaris_Control_Reporte_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# Mount static files for Frontend UI
STATIC_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static")
if os.path.exists(STATIC_DIR):
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

@app.get("/")
def read_root():
    if READ_ONLY_MODE:
        stakeholder_path = os.path.join(STATIC_DIR, "stakeholder.html")
        if os.path.exists(stakeholder_path):
            return FileResponse(stakeholder_path)
    index_path = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path)
    return JSONResponse({"message": "API de Granja Solar lista. Visite /docs para la documentación REST API."})

@app.get("/stakeholder")
def read_stakeholder_view():
    stakeholder_path = os.path.join(STATIC_DIR, "stakeholder.html")
    if os.path.exists(stakeholder_path):
        return FileResponse(stakeholder_path)
    raise HTTPException(status_code=404, detail="Vista de Stakeholders no encontrada.")

@app.get("/operator")
def read_operator_view():
    index_path = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path)
    raise HTTPException(status_code=404, detail="Vista de Operador no encontrada.")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("backend.main:app" if __package__ else "main:app", host="0.0.0.0", port=8000, reload=True)


